import streamlit as st
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

# ضبط إعدادات الصفحة
st.set_page_config(page_title="منصة تقديم البحوث الأكاديمية", layout="centered")

st.title("📄 نموذج تقديم وإعداد البحث الأكاديمي")
st.info("ملاحظة هامة: يُقبل البحث وفق نظام APA التوثيقي العالمي.")

with st.form("journal_form"):
    
    # ------------------ معلومات الباحثين ------------------
    st.header("معلومات الباحثين")
    
    st.subheader("الباحث الأول")
    b1_name = st.text_input("الاسم الكامل (الباحث الأول)")
    col1, col2 = st.columns(2)
    with col1:
        b1_univ = st.text_input("الجامعة (الباحث الأول)")
        b1_dept = st.text_input("القسم (الباحث الأول)")
        b1_email = st.text_input("البريد الإلكتروني (الباحث الأول)")
    with col2:
        b1_college = st.text_input("الكلية (الباحث الأول)")
        b1_country = st.text_input("البلد (الباحث الأول)")
        b1_phone = st.text_input("الهاتف (الباحث الأول)")

    st.markdown("---")
    st.subheader("الباحث الثاني (إن وجد)")
    b2_name = st.text_input("الاسم الكامل (الباحث الثاني)")
    col1, col2 = st.columns(2)
    with col1:
        b2_univ = st.text_input("الجامعة (الباحث الثاني)")
        b2_dept = st.text_input("القسم (الباحث الثاني)")
        b2_email = st.text_input("البريد الإلكتروني (الباحث الثاني)")
    with col2:
        b2_college = st.text_input("الكلية (الباحث الثاني)")
        b2_country = st.text_input("البلد (الباحث الثاني)")
        b2_phone = st.text_input("الهاتف (الباحث الثاني)")

    st.markdown("---")
    st.subheader("الباحث الثالث (إن وجد)")
    b3_name = st.text_input("الاسم الكامل (الباحث الثالث)")
    col1, col2 = st.columns(2)
    with col1:
        b3_univ = st.text_input("الجامعة (الباحث الثالث)")
        b3_dept = st.text_input("القسم (الباحث الثالث)")
        b3_email = st.text_input("البريد الإلكتروني (الباحث الثالث)")
    with col2:
        b3_college = st.text_input("الكلية (الباحث الثالث)")
        b3_country = st.text_input("البلد (الباحث الثالث)")
        b3_phone = st.text_input("الهاتف (الباحث الثالث)")

    st.markdown("---")
    # ------------------ القسم العربي ------------------
    st.header("بيانات البحث باللغة العربية")
    title_ar = st.text_input("عنوان البحث (باللغة العربية)")
    abstract_ar = st.text_area("ملخص البحث (باللغة العربية)", help="أن لا يتجاوز الملخص خارج القالب المخصص وبحجم خط 12", height=150)
    keywords_ar = st.text_input("الكلمات المفتاحية (باللغة العربية)", help="خمس كلمات مفتاحية تفصل بينها فاصلة")

    st.markdown("---")
    # ------------------ القسم الإنجليزي ------------------
    st.header("Research Details in English")
    title_en = st.text_input("Research Title (in English)")
    abstract_en = st.text_area("Abstract (in English)", help="Font size 12, fitting designated template space", height=150)
    keywords_en = st.text_input("Keywords (in English)", help="Five keywords separated by commas")

    st.markdown("---")
    # ------------------ متون البحث ------------------
    st.header("محتوى البحث")
    intro = st.text_area("المقدمة", height=200)
    body = st.text_area("متن البحث", height=300)
    references = st.text_area("المصادر والمراجع (وفق نظام APA)", height=200)

    submitted = st.form_submit_button("توليد ملف Word المنسق")

# عند الضغط على زر التوليد
if submitted:
    if not title_ar or not b1_name:
        st.error("يرجى إدخال عنوان البحث العربي واسم الباحث الأول على الأقل!")
    else:
        doc = Document()

        # ضبط الهوامش (1 إنش)
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        def add_heading(text, level=1, align=WD_ALIGN_PARAGRAPH.RIGHT):
            p = doc.add_paragraph()
            p.alignment = align
            run = p.add_run(text)
            run.font.name = 'Traditional Arabic'
            run.font.size = Pt(16 if level == 1 else 18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
            return p

        def add_body(text, font_size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_name='Traditional Arabic'):
            p = doc.add_paragraph()
            p.alignment = align
            run = p.add_run(text)
            run.font.name = font_name
            run.font.size = Pt(font_size)
            p.paragraph_format.line_spacing = 1.15
            return p

        # 1. عنوان البحث العربي
        p_t_ar = doc.add_paragraph()
        p_t_ar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t_ar = p_t_ar.add_run(title_ar)
        r_t_ar.font.name = 'Traditional Arabic'
        r_t_ar.font.size = Pt(20)
        r_t_ar.font.bold = True

        # 2. بيانات الباحثين
        authors_data = [
            (b1_name, b1_univ, b1_college, b1_dept, b1_country, b1_email, b1_phone),
            (b2_name, b2_univ, b2_college, b2_dept, b2_country, b2_email, b2_phone),
            (b3_name, b3_univ, b3_college, b3_dept, b3_country, b3_email, b3_phone)
        ]

        for idx, (name, univ, coll, dept, country, email, phone) in enumerate(authors_data, 1):
            if name:
                p_b = doc.add_paragraph()
                p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
                info_str = f"الباحث ({name})"
                details = [d for d in [dept, coll, univ, country] if d]
                if details:
                    info_str += " - " + "، ".join(details)
                if email:
                    info_str += f" | البريد: {email}"
                if phone:
                    info_str += f" | الهاتف: {phone}"
                
                r_b = p_b.add_run(info_str)
                r_b.font.name = 'Traditional Arabic'
                r_b.font.size = Pt(11)
                r_b.font.italic = True

        doc.add_paragraph()

        # 3. الملخص العربي والكلمات المفتاحية
        add_heading("ملخص البحث")
        add_body(abstract_ar, font_size=12) # خط 12 حسب طلبكم
        
        if keywords_ar:
            p_k = doc.add_paragraph()
            p_k.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_k1 = p_k.add_run("الكلمات المفتاحية: ")
            r_k1.font.bold = True
            r_k1.font.name = 'Traditional Arabic'
            r_k2 = p_k.add_run(keywords_ar)
            r_k2.font.name = 'Traditional Arabic'

        doc.add_paragraph()

        # 4. العنوان والملخص الإنجليزي
        if title_en:
            p_t_en = doc.add_paragraph()
            p_t_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_t_en = p_t_en.add_run(title_en)
            r_t_en.font.name = 'Times New Roman'
            r_t_en.font.size = Pt(18)
            r_t_en.font.bold = True

        if abstract_en:
            add_heading("Abstract", align=WD_ALIGN_PARAGRAPH.LEFT)
            add_body(abstract_en, font_size=12, align=WD_ALIGN_PARAGRAPH.LEFT, font_name='Times New Roman')

        if keywords_en:
            p_ke = doc.add_paragraph()
            p_ke.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_ke1 = p_ke.add_run("Keywords: ")
            r_ke1.font.bold = True
            r_ke1.font.name = 'Times New Roman'
            r_ke2 = p_ke.add_run(keywords_en)
            r_ke2.font.name = 'Times New Roman'

        doc.add_paragraph()

        # 5. متون البحث والمراجع
        if intro:
            add_heading("المقدمة")
            add_body(intro)

        if body:
            add_heading("متن البحث")
            add_body(body)

        if references:
